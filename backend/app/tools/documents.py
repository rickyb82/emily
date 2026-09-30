"""Estrazione testo da documenti + tool read_document per leggere file della conversazione/progetto."""
from __future__ import annotations

import csv
import io
from pathlib import Path

from sqlalchemy import select

from ..models import FileAsset, Share
from .registry import ToolContext, ToolResult, register

TEXT_MIMES = {"text/plain", "text/markdown", "text/csv", "application/json", "text/html", "application/xml"}
MAX_CHARS = 60_000


def extract_text(path: Path, mime: str) -> str | None:
    try:
        if mime in TEXT_MIMES or path.suffix.lower() in {".txt", ".md", ".json", ".csv", ".py", ".js", ".ts", ".html"}:
            return path.read_text(errors="replace")[:MAX_CHARS]
        if mime == "application/pdf" or path.suffix.lower() == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            parts = [f"--- pagina {i + 1} ---\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages)]
            return "\n".join(parts)[:MAX_CHARS]
        if path.suffix.lower() == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)[:MAX_CHARS]
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            out = io.StringIO()
            for ws in wb.worksheets:
                out.write(f"--- foglio {ws.title} ---\n")
                w = csv.writer(out)
                for row in ws.iter_rows(values_only=True):
                    w.writerow(["" if v is None else v for v in row])
            return out.getvalue()[:MAX_CHARS]
    except Exception as e:
        return f"[estrazione testo fallita: {e}]"
    return None


@register(
    "read_document",
    "Legge il contenuto testuale di un documento caricato dall'utente (PDF, Word, Excel, testo) dato il suo id. "
    "Usalo quando l'utente fa riferimento a un file allegato alla chat o al progetto.",
    {"type": "object", "properties": {"file_id": {"type": "string", "description": "Id del file"}},
     "required": ["file_id"]},
)
async def read_document(args: dict, ctx: ToolContext) -> ToolResult:
    asset = await ctx.db.get(FileAsset, str(args.get("file_id", "")))
    if not asset or asset.family_id != ctx.user.family_id:
        return ToolResult("File non trovato", is_error=True)
    ok = asset.owner_id == ctx.user.id or (ctx.project_id and asset.project_id == ctx.project_id)
    if not ok:
        share = (await ctx.db.execute(select(Share).where(
            Share.resource_type == "file", Share.resource_id == asset.id,
            (Share.target_user_id == ctx.user.id) | (Share.target_user_id.is_(None))))).scalar_one_or_none()
        ok = share is not None
    if not ok:
        return ToolResult("Non hai accesso a questo file", is_error=True)
    if asset.extracted_text is None:
        asset.extracted_text = extract_text(Path(asset.path), asset.mime) or ""
        await ctx.db.commit()
    return ToolResult(asset.extracted_text or "(nessun testo estraibile)")
